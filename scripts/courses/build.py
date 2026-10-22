#!/usr/bin/env python3
"""Build the facilitator courses in courses/ into teachable outputs.

For every course folder that has a slides.md:
  * courses/build/<course>/<deck>.pptx   one PowerPoint per day, speaker notes = the short script
  * courses/build/<course>/<course>-phone-pack.html   every Markdown file of the course (plus the
    shared toolkit and forms) in one offline HTML page a facilitator can read on a phone
And once for all courses:
  * courses/media-manifest.json   every image / animation / poster artwork, with its brief, status
    (NEW for Codex, REUSE of an existing repo file) and the path Codex should write to

WHY ONE SOURCE: the day files, posters.md, slides.md and media.md are the course. Decks, phone
packs and the Codex manifest are derived from them, so a correction is made once in Markdown and
re-built, never patched in a .pptx by hand.

Media are placed on slides when the file exists (REUSE path, or public/course-media/<course>/<ID>.*
once Codex has made it); otherwise the slide carries a labelled placeholder with the brief, so a
deck is teachable today and fills itself in as art arrives. Clips are NOT embedded by default (a
day's clips would make a 60 MB deck); the slide shows the clip's still and names the file to play
from the tablet. Pass --embed-video to embed them.

Usage:
  pip install python-pptx markdown
  python3 scripts/courses/build.py                 # all courses
  python3 scripts/courses/build.py farmer-5day     # one course
  python3 scripts/courses/build.py --check         # validate references only, write nothing
"""
from __future__ import annotations

import argparse
import html
import json
import re
import sys
from dataclasses import dataclass, field
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
COURSES = ROOT / "courses"
BUILD = COURSES / "build"
PUBLIC = ROOT / "public"

PREFIX_TO_COURSE = {"F": "farmer-5day", "M": "mentor-3day", "T": "teach-the-teachers",
                    "R": "farmer-money", "A": "ai-literacy"}
COURSE_TITLES = {
    "farmer-5day": "Grow Food, Grow Strong",
    "mentor-3day": "Walking With Farmers",
    "teach-the-teachers": "Teach the Teachers",
    "farmer-money": "Farmer Money",
    "ai-literacy": "AI Literacy",
}

# ImbewuField palette (CLAUDE.md)
FOREST, MID, LEAF, OCHRE, WATER = "1F4D2B", "2E6B3A", "A8D88A", "C07A1E", "235E86"
PAPER, CARD, INK, HAIR = "F7F2E9", "FBF6EC", "20190F", "E2D8C4"
# Portable fonts: facilitator laptops won't have Newsreader / Public Sans installed.
HEAD_FONT, BODY_FONT = "Georgia", "Arial"

ID_RE = re.compile(r"\b([FMTRA])-(?:IMG|ANI|PHO)-\d+\b|\b([FMTRA])-P\d{2}-ART\b")
ANY_MEDIA_RE = re.compile(r"\b[FMTRA]-(?:IMG|ANI|PHO)-\d+\b|\b[FMTRA]-P\d{2}-ART\b")
POSTER_RE = re.compile(r"\b[FMTRA]-P\d{2}\b(?!-ART)")


# ── parsing ──────────────────────────────────────────────────────────────────────────────────

@dataclass
class Slide:
    id: str
    title: str
    type: str = "content"
    track: str = ""
    screen: list[str] = field(default_factory=list)
    media: list[str] = field(default_factory=list)
    posters: list[str] = field(default_factory=list)
    notes: str = ""


@dataclass
class Deck:
    id: str
    title: str
    slides: list[Slide] = field(default_factory=list)


def split_ids(value: str) -> list[str]:
    return [v.strip() for v in re.split(r"[,;]", value) if v.strip() and v.strip() not in ("—", "-")]


def parse_slides(path: Path) -> list[Deck]:
    decks: list[Deck] = []
    slide: Slide | None = None
    mode = None
    for raw in path.read_text(encoding="utf-8").splitlines():
        line = raw.rstrip()
        m = re.match(r"^# Deck (\S+)\s*[·—-]\s*(.+)$", line)
        if m:
            decks.append(Deck(m.group(1), m.group(2).strip()))
            slide, mode = None, None
            continue
        m = re.match(r"^## (\S+)\s*·\s*(.+)$", line)
        if m and decks:
            slide = Slide(m.group(1), m.group(2).strip())
            decks[-1].slides.append(slide)
            mode = None
            continue
        if slide is None:
            continue
        m = re.match(r"^(Type|Track|Media|Poster|Screen|Notes):\s*(.*)$", line)
        if m and mode != "notes":
            key, val = m.group(1), m.group(2).strip()
            if key == "Type":
                slide.type = val.lower()
            elif key == "Track":
                slide.track = val.lower()
            elif key == "Media":
                slide.media = split_ids(val)
            elif key == "Poster":
                slide.posters = split_ids(val)
            elif key == "Screen":
                mode = "screen"
            elif key == "Notes":
                mode = "notes"
                if val:
                    slide.notes += val + "\n"
            continue
        if mode == "screen" and line.startswith("- "):
            slide.screen.append(line[2:].strip())
        elif mode == "notes":
            slide.notes += line + "\n"
    for d in decks:
        for s in d.slides:
            s.notes = s.notes.strip()
    return decks


def parse_blocks(path: Path, header_re: str) -> list[dict]:
    """Parse `## <ID> · <title>` blocks with `- Key: value` lines and a free body."""
    if not path.exists():
        return []
    items: list[dict] = []
    cur: dict | None = None
    for line in path.read_text(encoding="utf-8").splitlines():
        m = re.match(r"^## (" + header_re + r")\s*·\s*(.+)$", line)
        if m:
            cur = {"id": m.group(1), "title": m.group(2).strip(), "fields": {}, "body": []}
            items.append(cur)
            continue
        if line.startswith("## ") or line.startswith("# "):
            cur = None
            continue
        if cur is None:
            continue
        m = re.match(r"^- ([A-Za-z ]+):\s*(.*)$", line)
        if m and not cur["body"]:
            cur["fields"][m.group(1).strip().lower()] = m.group(2).strip()
        else:
            cur["body"].append(line)
    for it in items:
        it["body"] = re.sub(r"\n-{3,}\s*$", "", "\n".join(it["body"]).strip()).strip()
    return items


def section(body: str, label: str) -> str:
    """Pull the text after a bold **Label…:** up to the next bold label."""
    m = re.search(r"\*\*" + re.escape(label) + r"[^*]*\*\*\s*(.*?)(?=\n\*\*[A-Z][^*]*:\*\*|\Z)", body, re.S)
    return m.group(1).strip() if m else ""


# ── media resolution ─────────────────────────────────────────────────────────────────────────

def course_of(media_id: str) -> str:
    return PREFIX_TO_COURSE[media_id[0]]


def target_path(media_id: str, typ: str) -> str:
    ext = "mp4" if typ == "animation" else "jpg"
    return f"public/course-media/{course_of(media_id)}/{media_id}.{ext}"


def build_catalog(course_dirs: list[Path]) -> dict[str, dict]:
    catalog: dict[str, dict] = {}
    for cdir in course_dirs:
        course = cdir.name
        for it in parse_blocks(cdir / "media.md", r"[FMTRA]-(?:IMG|ANI|PHO)-\d+"):
            f = it["fields"]
            typ = f.get("type", "image").lower()
            status = f.get("status", "NEW")
            reuse = status.split(None, 1)[1].strip() if status.upper().startswith("REUSE") and " " in status else ""
            brief = re.sub(r"^\*\*Brief:\*\*\s*", "", it["body"]).strip()
            catalog[it["id"]] = {
                "id": it["id"], "course": course, "kind": typ, "title": it["title"],
                "status": "REUSE" if reuse else "NEW", "reuse_path": reuse or None,
                "target": None if reuse else target_path(it["id"], typ),
                "used_in": split_ids(f.get("used in", "")), "length": f.get("length", ""),
                "brief": brief,
            }
        for it in parse_blocks(cdir / "posters.md", r"[FMTRA]-P\d{2}"):
            f = it["fields"]
            art_id = f.get("art") or f"{it['id']}-ART"
            poster = {
                "id": it["id"], "course": course, "kind": "poster", "title": it["title"],
                "size": f.get("size", "A1 portrait"), "used_in": split_ids(f.get("used in", "")),
                "art": art_id, "headline": section(it["body"], "Headline"),
                "labels": section(it["body"], "Labels"), "layout": section(it["body"], "Layout"),
                "teach_it": section(it["body"], "Teach it"), "status": "NEW",
                "target": f"public/course-media/{course}/posters/{it['id']}.pdf",
                "phone_target": f"public/course-media/{course}/posters/{it['id']}.jpg",
            }
            catalog[it["id"]] = poster
            catalog[art_id] = {
                "id": art_id, "course": course, "kind": "poster-art", "title": f"Artwork for {it['id']} · {it['title']}",
                "status": "NEW", "reuse_path": None, "target": target_path(art_id, "image"),
                "used_in": [it["id"]] + poster["used_in"], "length": "still",
                "brief": section(it["body"], "Art brief"),
            }
    return catalog


def existing_file(item: dict) -> Path | None:
    for rel in (item.get("reuse_path"), item.get("target")):
        if rel:
            p = ROOT / rel
            if p.exists():
                return p
    return None


def still_for(item: dict) -> Path | None:
    """An image to put on the slide: the image itself, or a clip's poster frame."""
    f = existing_file(item)
    if f is None:
        return None
    if f.suffix.lower() in (".jpg", ".jpeg", ".png", ".webp"):
        return f if f.suffix.lower() != ".webp" else None
    if f.suffix.lower() == ".mp4":
        for cand in (f.parent / "posters" / (f.stem + ".jpg"), f.with_suffix(".jpg")):
            if cand.exists():
                return cand
    return None


# ── per-course look (courses/<course>/style.md) ──────────────────────────────────────────────

def _code_after(md: str, heading: str) -> str:
    m = re.search(r"^##\s+" + re.escape(heading) + r".*?\n```[a-z]*\n(.*?)\n```", md, re.S | re.M)
    return m.group(1).strip() if m else ""


def load_style(course: str) -> dict:
    """Deck theme + generator anchors for a course. Falls back to the ImbewuField palette."""
    p = COURSES / course / "style.md"
    md = p.read_text(encoding="utf-8") if p.exists() else ""
    theme = {"bg": PAPER, "ink": INK, "accent": OCHRE, "accent2": WATER, "section_bg": FOREST,
             "section_ink": PAPER, "fallback_head": HEAD_FONT, "fallback_body": BODY_FONT}
    raw = _code_after(md, "Deck theme")
    if raw:
        theme.update(json.loads(raw))
    title = re.search(r"^# .*?[—-]\s*\"?(.+?)\"?\s*$", md, re.M)
    return {"theme": theme, "look": title.group(1) if title else "",
            "image_anchor": _code_after(md, "Image anchor"),
            "poster_anchor": _code_after(md, "Poster anchor"),
            "animation_anchor": _code_after(md, "Animation anchor")}


def load_i18n(course: str, lang: str) -> dict:
    p = COURSES / course / "i18n" / f"{lang}.json"
    return json.loads(p.read_text(encoding="utf-8")) if p.exists() else {}


# ── validation ───────────────────────────────────────────────────────────────────────────────

def validate(course_dirs: list[Path], catalog: dict) -> list[str]:
    problems: list[str] = []
    for cdir in course_dirs:
        decks = parse_slides(cdir / "slides.md") if (cdir / "slides.md").exists() else []
        if not decks:
            problems.append(f"{cdir.name}: no decks in slides.md")
        seen: set[str] = set()
        for d in decks:
            for s in d.slides:
                if s.id in seen:
                    problems.append(f"{cdir.name}: duplicate slide id {s.id}")
                seen.add(s.id)
                if len(s.screen) > 3:
                    problems.append(f"{s.id}: {len(s.screen)} screen bullets (max 3)")
                for mid in s.media:
                    if mid not in catalog:
                        problems.append(f"{s.id}: media {mid} is not defined in any media.md / posters.md")
                for pid in s.posters:
                    if pid not in catalog:
                        problems.append(f"{s.id}: poster {pid} is not defined in any posters.md")
        # IDs mentioned anywhere in the course's Markdown must resolve too
        for md in sorted(cdir.glob("*.md")):
            text = md.read_text(encoding="utf-8")
            for mid in set(ANY_MEDIA_RE.findall(text)):
                if mid not in catalog:
                    problems.append(f"{cdir.name}/{md.name}: mentions {mid}, which is not defined")
            for pid in set(POSTER_RE.findall(text)):
                if pid not in catalog:
                    problems.append(f"{cdir.name}/{md.name}: mentions poster {pid}, which is not defined")
        for it in catalog.values():
            if it["course"] == cdir.name and it.get("reuse_path") and not (ROOT / it["reuse_path"]).exists():
                problems.append(f"{it['id']}: REUSE path does not exist: {it['reuse_path']}")
    return problems


# ── PowerPoint ───────────────────────────────────────────────────────────────────────────────

def _mix(a: str, b: str, t: float) -> str:
    """Blend two hex colours: t=0 → a, t=1 → b."""
    ca = [int(a[i:i + 2], 16) for i in (0, 2, 4)]
    cb = [int(b[i:i + 2], 16) for i in (0, 2, 4)]
    return "".join(f"{round(x + (y - x) * t):02X}" for x, y in zip(ca, cb))


def _deck_kicker(deck: Deck, lang: str) -> str:
    """'Day 1 — Why we farm this way' → 'DAY 1' (isiZulu: 'USUKU 1')."""
    head = re.split(r"\s+[—-]\s+", deck.title, maxsplit=1)[0].strip()
    if lang == "zu":
        head = re.sub(r"^Day\b", "Usuku", head)
    return head.upper()


# Per-course decoration (style.md "motif"): every course shares one layout grid, and the motif
# decides frames, bullet markers and the section-slide ornament.
MOTIFS = {
    "linocut": {"marker": "■", "round": False, "frame_w": 4.0},
    "photo":   {"marker": "—", "round": False, "frame_w": 0.0},
    "chalk":   {"marker": "›", "round": False, "frame_w": 0.0},
    "clay":    {"marker": "●", "round": True,  "frame_w": 0.0},
    "comic":   {"marker": "■", "round": False, "frame_w": 4.5},
}


def build_deck(course: str, deck: Deck, catalog: dict, out: Path, embed_video: bool, lang: str = "en") -> None:
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
    from pptx.oxml.ns import qn
    from pptx.util import Emu, Pt

    def rgb(h: str) -> RGBColor:
        return RGBColor.from_string(h)

    style = load_style(course)
    T = style["theme"]
    # Shadow the house palette with this course's own look (see style.md).
    PAPER, INK, ACC, ACC2 = T["bg"], T["ink"], T["accent"], T["accent2"]
    SEC_BG, SEC_INK = T["section_bg"], T["section_ink"]
    MUTED = _mix(INK, PAPER, 0.45)
    CARD = T.get("card", _mix(PAPER, "FFFFFF", 0.55))
    HEAD_FONT, BODY_FONT = T["fallback_head"], T["fallback_body"]
    MOTIF = T.get("motif", "linocut")
    MO = MOTIFS.get(MOTIF, MOTIFS["linocut"])
    tr = (load_i18n(course, lang).get("slides") or {}) if lang != "en" else {}
    kicker = _deck_kicker(deck, lang)

    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(12192000), Emu(6858000)  # 16:9, 13.333 × 7.5 in
    W, H = prs.slide_width, prs.slide_height
    blank = prs.slide_layouts[6]
    IN = 914400
    M = 0.65 * IN  # outer margin

    def alpha(shape, opacity: float) -> None:
        fill = shape.fill._xPr.find(qn("a:solidFill"))
        if fill is not None and len(fill):
            el = fill[0].makeelement(qn("a:alpha"), {"val": str(int(opacity * 100000))})
            fill[0].append(el)

    def box(slide, x, y, w, h, fill=None, line=None, lw=1.25, shape=None, rot=0.0, radius=None):
        if shape is None:
            shape = MSO_SHAPE.ROUNDED_RECTANGLE if MO["round"] else MSO_SHAPE.RECTANGLE
        s = slide.shapes.add_shape(shape, Emu(int(x)), Emu(int(y)), Emu(int(w)), Emu(int(h)))
        if fill:
            s.fill.solid(); s.fill.fore_color.rgb = rgb(fill)
        else:
            s.fill.background()
        if line:
            s.line.color.rgb = rgb(line); s.line.width = Pt(lw)
        else:
            s.line.fill.background()
        if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
            s.adjustments[0] = radius if radius is not None else 0.12
        if rot:
            s.rotation = rot
        s.shadow.inherit = False
        return s

    def text(slide, x, y, w, h, runs, size=24, color=INK, font=BODY_FONT, bold=False,
             align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, marker=None, marker_color=None, spacing=8,
             line_spacing=None):
        tb = slide.shapes.add_textbox(Emu(int(x)), Emu(int(y)), Emu(int(w)), Emu(int(h)))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = anchor
        tf.margin_left = tf.margin_right = Emu(int(0.05 * IN))
        tf.margin_top = tf.margin_bottom = Emu(0)
        for i, r in enumerate(runs if isinstance(runs, list) else [runs]):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = align
            p.space_after = Pt(spacing)
            if line_spacing:
                p.line_spacing = line_spacing
            if marker:
                hang = int(size * 1.25 * 12700)  # hanging indent so wrapped lines align with the text
                pPr = p._p.get_or_add_pPr()
                pPr.set("marL", str(hang)); pPr.set("indent", str(-hang))
                m = p.add_run()
                m.text = marker + "  "
                m.font.size = Pt(size); m.font.bold = True; m.font.name = font
                m.font.color.rgb = rgb(marker_color or ACC)
            run = p.add_run()
            run.text = r
            run.font.size = Pt(size); run.font.bold = bold; run.font.name = font
            run.font.color.rgb = rgb(color)
        return tb

    def pill(slide, x, y, label, fill, ink, size=12, w=None, outline=False):
        w = w or (0.32 * IN + len(label) * size * 0.62 * IN / 72)
        h = size * 2.1 * IN / 72
        if MOTIF == "comic" and not outline:
            box(slide, x + 0.05 * IN, y + 0.05 * IN, w, h, fill=INK)
            b = box(slide, x, y, w, h, fill=fill, line=INK, lw=2.0)
        else:
            b = box(slide, x, y, w, h, fill=None if outline else fill, line=fill if outline else None,
                    lw=1.25, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.5)
        text(slide, x, y, w, h, label, size=size, color=fill if outline else ink, bold=True,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        return w

    # ── frames: one per motif ────────────────────────────────────────────────────────────────
    def frame(slide, x, y, w, h, dark=False):
        """Draw the course's picture frame; return the inner box the picture sits in."""
        ink = SEC_INK if dark else INK
        if MOTIF == "linocut":
            box(slide, x, y, w, h, fill=CARD, line=ink, lw=MO["frame_w"])
            return x + 0.12 * IN, y + 0.12 * IN, w - 0.24 * IN, h - 0.24 * IN
        if MOTIF == "photo":
            sh = box(slide, x + 0.09 * IN, y + 0.11 * IN, w, h, fill=_mix(INK, PAPER, 0.75))
            alpha(sh, 0.55)
            box(slide, x, y, w, h, fill="FFFFFF", line=_mix(INK, PAPER, 0.8), lw=0.75)
            return x + 0.16 * IN, y + 0.16 * IN, w - 0.32 * IN, h - 0.56 * IN
        if MOTIF == "chalk":
            box(slide, x, y, w, h, fill="FBF8F1", rot=-1.2)
            for tx, rot in ((x + 0.25 * IN, -18), (x + w - 1.15 * IN, 14)):
                tape = box(slide, tx, y - 0.14 * IN, 0.9 * IN, 0.3 * IN, fill="F3E6A6", rot=rot)
                alpha(tape, 0.8)
            return x + 0.2 * IN, y + 0.2 * IN, w - 0.4 * IN, h - 0.4 * IN
        if MOTIF == "clay":
            box(slide, x, y, w, h, fill=CARD, line=_mix(ACC, PAPER, 0.4), lw=3.0, radius=0.08)
            return x + 0.18 * IN, y + 0.18 * IN, w - 0.36 * IN, h - 0.36 * IN
        # comic: hard offset shadow + heavy ink border
        box(slide, x + 0.13 * IN, y + 0.13 * IN, w, h, fill=INK)
        box(slide, x, y, w, h, fill="FFFFFF", line=INK, lw=MO["frame_w"])
        return x + 0.12 * IN, y + 0.12 * IN, w - 0.24 * IN, h - 0.24 * IN

    def place_media(slide, mid: str, x, y, w, h, dark=False):
        item = catalog.get(mid)
        if item is None:
            return
        ix, iy, iw, ih = frame(slide, x, y, w, h, dark)
        f = existing_file(item)
        if embed_video and f is not None and f.suffix.lower() == ".mp4":
            poster = still_for(item)
            slide.shapes.add_movie(str(f), Emu(int(ix)), Emu(int(iy)), Emu(int(iw)), Emu(int(ih)),
                                   poster_frame_image=str(poster) if poster else None, mime_type="video/mp4")
            return
        still = still_for(item)
        if still is not None:
            from PIL import Image
            with Image.open(still) as im:
                pw0, ph0 = im.size
            scale = min(iw / pw0, ih / ph0)
            pw, ph = pw0 * scale, ph0 * scale
            slide.shapes.add_picture(str(still), Emu(int(ix + (iw - pw) / 2)), Emu(int(iy + (ih - ph) / 2)),
                                     Emu(int(pw)), Emu(int(ph)))
            if item["kind"] == "animation":
                rel = (item.get("reuse_path") or item.get("target") or "")
                tag = box(slide, ix, iy + ih - 0.45 * IN, iw, 0.45 * IN, fill=INK)
                alpha(tag, 0.8)
                text(slide, ix + 0.12 * IN, iy + ih - 0.45 * IN, iw - 0.24 * IN, 0.45 * IN,
                     f"▶  Play clip · {Path(rel).name}", size=12, color="FFFFFF", anchor=MSO_ANCHOR.MIDDLE)
            return
        # Placeholder until the art lands: a quiet, labelled frame (the full brief is in the notes).
        kind = {"animation": "CLIP", "poster-art": "POSTER ART", "photo": "PHOTO"}.get(item["kind"], "PICTURE")
        if lang == "zu":
            kind = {"CLIP": "IVIDIYO", "PHOTO": "ISITHOMBE", "PICTURE": "ISITHOMBE"}.get(kind, kind)
        mark = box(slide, ix + iw / 2 - 0.45 * IN, iy + ih / 2 - 1.05 * IN, 0.9 * IN, 0.9 * IN,
                   fill=None, line=_mix(ACC, CARD, 0.35), lw=2.0,
                   shape=MSO_SHAPE.OVAL if kind in ("CLIP", "IVIDIYO") else
                   (MSO_SHAPE.ROUNDED_RECTANGLE if MO["round"] else MSO_SHAPE.RECTANGLE))
        text(slide, ix + iw / 2 - 0.45 * IN, iy + ih / 2 - 1.05 * IN, 0.9 * IN, 0.9 * IN,
             "▶" if kind in ("CLIP", "IVIDIYO") else "✦", size=26, color=_mix(ACC, CARD, 0.2),
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        text(slide, ix + 0.3 * IN, iy + ih / 2 + 0.05 * IN, iw - 0.6 * IN, 0.9 * IN, item["title"], size=20,
             color=_mix(INK, CARD, 0.15), font=HEAD_FONT, bold=True, align=PP_ALIGN.CENTER)
        text(slide, ix + 0.3 * IN, iy + ih - 0.5 * IN, iw - 0.6 * IN, 0.3 * IN, f"{kind} · {mid}", size=10,
             color=MUTED, align=PP_ALIGN.CENTER)

    # ── ornaments ────────────────────────────────────────────────────────────────────────────
    def ornament_section(slide):
        if MOTIF == "linocut":
            box(slide, 0.22 * IN, 0.22 * IN, W - 0.44 * IN, H - 0.44 * IN, line=SEC_INK, lw=2.5)
            n, tw = 22, (W - 0.44 * IN) / 22
            for i in range(n):
                t = box(slide, 0.22 * IN + i * tw, H - 0.62 * IN, tw, 0.4 * IN, fill=ACC,
                        shape=MSO_SHAPE.ISOSCELES_TRIANGLE)
        elif MOTIF == "photo":
            for yy in (0.5 * IN, 0.58 * IN, H - 0.62 * IN, H - 0.54 * IN):
                box(slide, M, yy, W - 2 * M, Emu(int(0.012 * IN)), fill=ACC)
        elif MOTIF == "chalk":
            smudge = box(slide, -0.6 * IN, H - 2.1 * IN, 5.5 * IN, 2.6 * IN, fill=SEC_INK, shape=MSO_SHAPE.OVAL)
            alpha(smudge, 0.05)
            smudge = box(slide, W - 4.5 * IN, -0.8 * IN, 5.2 * IN, 2.4 * IN, fill=SEC_INK, shape=MSO_SHAPE.OVAL)
            alpha(smudge, 0.04)
        elif MOTIF == "clay":
            for cx, cy, d, col, a in ((-1.1, -1.0, 3.4, ACC, 0.9), (W / IN - 1.5, 2.4, 2.9, ACC2, 0.85),
                                      (W / IN - 3.2, -0.6, 1.3, SEC_INK, 0.18), (5.6, -0.45, 0.95, ACC2, 0.7)):
                blob = box(slide, cx * IN, cy * IN, d * IN, d * 0.92 * IN, fill=col, shape=MSO_SHAPE.OVAL)
                alpha(blob, a)
        else:  # comic: halftone corner + burst
            for r in range(6):
                for c in range(10 - r):
                    d = (0.16 - 0.022 * r) * IN
                    dot = box(slide, W - (c + 1) * 0.3 * IN, (0.25 + r * 0.3) * IN, d, d, fill=ACC, shape=MSO_SHAPE.OVAL)
                    alpha(dot, 0.6)

    def ornament_content(slide, activity: bool):
        if MOTIF == "linocut":
            box(slide, 0.22 * IN, 0.22 * IN, W - 0.44 * IN, H - 0.44 * IN, line=INK, lw=1.5)
        if activity:
            band = box(slide, 0, 0, 0.3 * IN, H, fill=ACC, shape=MSO_SHAPE.RECTANGLE)
            if MOTIF == "comic":
                box(slide, 0.3 * IN, 0, Emu(int(0.05 * IN)), H, fill=INK, shape=MSO_SHAPE.RECTANGLE)

    def title_rule(slide, x, y, w):
        if MOTIF == "chalk":
            box(slide, x, y, 2.6 * IN, 0.07 * IN, fill=ACC, rot=-0.8)
            box(slide, x + 0.3 * IN, y + 0.1 * IN, 1.8 * IN, 0.04 * IN, fill=ACC, rot=0.6)
        elif MOTIF == "clay":
            box(slide, x, y, 1.4 * IN, 0.12 * IN, fill=ACC, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.5)
        elif MOTIF == "comic":
            box(slide, x, y, 1.6 * IN, 0.12 * IN, fill=ACC2, shape=MSO_SHAPE.RECTANGLE)
        elif MOTIF == "photo":
            box(slide, x, y, w, Emu(int(0.012 * IN)), fill=_mix(INK, PAPER, 0.7))
        else:
            box(slide, x, y, 1.2 * IN, 0.1 * IN, fill=ACC)

    def lines(s_: str, w, size, bold=False) -> int:
        cpl = max(1, int((w / IN) * 72 / (size * (0.66 if bold else 0.56))))
        if max((len(wd) for wd in s_.split()), default=0) > cpl:
            return 99  # a single word would break mid-word
        n, cur = 1, 0
        for word in s_.split():
            if cur and cur + 1 + len(word) > cpl:
                n, cur = n + 1, len(word)
            else:
                cur += (1 if cur else 0) + len(word)
        return n

    def number_badge(slide, x, y, n: int):
        d = 0.62 * IN
        if MOTIF == "comic":
            box(slide, x + 0.05 * IN, y + 0.05 * IN, d, d, fill=INK, shape=MSO_SHAPE.OVAL)
            box(slide, x, y, d, d, fill=ACC, line=INK, lw=2.0, shape=MSO_SHAPE.OVAL)
        else:
            box(slide, x, y, d, d, fill=ACC, shape=MSO_SHAPE.OVAL)
        text(slide, x, y, d, d, str(n), size=22, color="FFFFFF", font=HEAD_FONT, bold=True,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

    total = len(deck.slides)
    for idx, s in enumerate(deck.slides, start=1):
        t = tr.get(s.id) or {}
        title = t.get("title") or s.title
        screen = t.get("screen") if t.get("screen") and len(t["screen"]) == len(s.screen) else s.screen
        # "[APP] Record visit" / "[PAPER] The visit form": the track shows as a pill instead.
        mt = re.match(r"^\[(APP|PAPER)\]\s*", title)
        track = (mt.group(1) if mt else ("APP" if s.track == "app" else "")).upper()
        title = title[mt.end():] if mt else title
        slide = prs.slides.add_slide(blank)
        section = s.type == "section"
        activity = s.type == "activity"
        bg = slide.background.fill
        bg.solid(); bg.fore_color.rgb = rgb(SEC_BG if section else PAPER)
        ink = SEC_INK if section else INK
        muted = _mix(SEC_INK, SEC_BG, 0.4) if section else MUTED

        if section:
            ornament_section(slide)
        else:
            ornament_content(slide, activity)

        # footer: course · poster to switch to · page
        fy = H - 0.62 * IN
        if MOTIF == "linocut" and section:
            fy = H - 1.05 * IN
        fx = M
        if s.posters:
            label = ("IPHOSTA " if lang == "zu" else "POSTER ") + ", ".join(s.posters)
            fx += pill(slide, fx, fy - 0.02 * IN, label, ACC2 if not section else SEC_INK, PAPER, size=10,
                       outline=True) + 0.15 * IN
        text(slide, fx, fy, 6 * IN, 0.3 * IN, f"{COURSE_TITLES.get(course, course)}  ·  {s.id}", size=10,
             color=muted, anchor=MSO_ANCHOR.MIDDLE)
        text(slide, W - M - 2 * IN, fy, 2 * IN, 0.3 * IN, f"{idx} / {total}", size=10, color=muted,
             align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE)
        if track:
            label = {"APP": "APP", "PAPER": "IPHEPHA" if lang == "zu" else "PAPER"}[track]
            pill(slide, W - M - 1.25 * IN, 0.45 * IN, label, ACC2, "FFFFFF", size=12, w=1.25 * IN,
                 outline=(track == "PAPER"))

        if section:
            tx, tw = M + 0.25 * IN, (7.0 if s.media else 11.5) * IN
            text(slide, tx, 1.05 * IN, tw, 0.4 * IN, kicker, size=16, color=_mix(SEC_INK, SEC_BG, 0.25),
                 bold=True)
            ssize = next((z for z, mx in ((54, 2), (48, 2), (42, 2), (38, 3), (34, 3), (30, 3))
                          if lines(title, tw, z, True) <= mx), 30)
            text(slide, tx, 1.5 * IN, tw, 2.3 * IN, title, size=ssize, color=SEC_INK, font=HEAD_FONT, bold=True,
                 anchor=MSO_ANCHOR.BOTTOM, line_spacing=0.95)
            if MOTIF == "clay":
                box(slide, tx, 3.98 * IN, 1.6 * IN, 0.14 * IN, fill=ACC, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.5)
            elif MOTIF == "chalk":
                box(slide, tx, 3.98 * IN, 2.8 * IN, 0.07 * IN, fill=ACC, rot=-0.8)
            else:
                box(slide, tx, 3.98 * IN, 1.6 * IN, 0.1 * IN, fill=ACC if MOTIF != "comic" else SEC_INK)
            text(slide, tx, 4.3 * IN, tw, 1.8 * IN, screen, size=24, color=_mix(SEC_INK, SEC_BG, 0.15))
            if s.media:
                place_media(slide, s.media[0], 8.15 * IN, 1.0 * IN, 4.5 * IN, 4.6 * IN, dark=True)
        else:
            x0 = M + (0.15 * IN if activity else 0)
            ty = 0.5 * IN
            if activity:
                lab = "SEBENZANI" if lang == "zu" else "YOUR TURN"
                pill(slide, x0, ty, lab, ACC, "FFFFFF", size=13)
                ty += 0.55 * IN
            else:
                text(slide, x0, ty, 9 * IN, 0.35 * IN, kicker, size=13, color=ACC, bold=True)
                ty += 0.38 * IN
            right = W - M - (1.4 * IN if track else 0)
            tsize = next((z for z, mx in ((40, 1), (36, 1), (32, 1), (32, 2), (28, 2)) if lines(title, right - x0, z, True) <= mx), 28)
            tl = min(lines(title, right - x0, tsize, True), 3)
            th = tl * tsize * 1.05 * IN / 72
            text(slide, x0, ty, right - x0, th + 0.1 * IN, title, size=tsize, color=INK, font=HEAD_FONT, bold=True,
                 anchor=MSO_ANCHOR.TOP, line_spacing=0.95)
            title_rule(slide, x0 + 0.05 * IN, ty + th + 0.22 * IN, right - x0)
            by = ty + th + 0.62 * IN
            avail = H - by - 0.95 * IN
            has_media = bool(s.media)
            cw = (5.55 * IN if has_media else W - M - x0) - (0.15 * IN if activity else 0)
            if activity:
                for size in ((26, 24, 22, 20) if has_media else (30, 28, 26, 24)):
                    rows = [max(0.78 * IN, lines(ln, cw - 0.85 * IN, size, True) * size * 1.12 * IN / 72 + 0.28 * IN)
                            for ln in screen]
                    if sum(rows) <= avail:
                        break
                ry = by
                for n, (line, rh) in enumerate(zip(screen, rows), start=1):
                    number_badge(slide, x0, ry, n)
                    text(slide, x0 + 0.85 * IN, ry + 0.05 * IN, cw - 0.85 * IN, rh - 0.1 * IN, line, size=size,
                         color=INK, bold=True, line_spacing=0.95)
                    ry += rh
            else:
                for size in ((26, 24, 22, 20) if has_media else (32, 30, 28, 24)):
                    need = sum(lines(b, cw - size * 1.25 * IN / 72, size) * size * 1.2 * IN / 72 + 18 * IN / 72
                               for b in screen)
                    if need <= avail:
                        break
                text(slide, x0, by, cw, avail, screen, size=size, color=INK, marker=MO["marker"], spacing=18)
            if has_media:
                fx0 = 6.55 * IN
                place_media(slide, s.media[0], fx0, by - 0.1 * IN, W - M - fx0, avail + 0.05 * IN)

        notes = s.notes
        extra = []
        if s.posters:
            extra.append("Poster: " + ", ".join(s.posters))
        if len(s.media) > 1:
            extra.append("Also: " + ", ".join(s.media[1:]))
        for m in s.media:
            it = catalog.get(m)
            if it is None:
                continue
            if it.get("kind") == "animation":
                extra.append(f"Clip {m}: {it.get('reuse_path') or it.get('target')}")
            if existing_file(it) is None and it.get("brief"):
                brief = re.sub(r"\s+", " ", it["brief"]).strip()
                extra.append(f"Picture {m} (coming) — {it['title']}: {brief}")
        slide.notes_slide.notes_text_frame.text = (notes + ("\n\n" + "\n".join(extra) if extra else "")).strip()

    out.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(out))


# ── phone pack ───────────────────────────────────────────────────────────────────────────────

PACK_CSS = """
:root{--forest:#1F4D2B;--mid:#2E6B3A;--leaf:#A8D88A;--ochre:#C07A1E;--water:#235E86;--paper:#F7F2E9;
--card:#FBF6EC;--ink:#20190F;--hair:#E2D8C4;--muted:#5b5446}
@media (prefers-color-scheme:dark){:root:not([data-theme="light"]){--paper:#15140f;--card:#1d1b15;--ink:#efe8da;
--hair:#39342a;--muted:#b5ab98;--forest:#A8D88A;--mid:#8fc775}}
:root[data-theme="dark"]{--paper:#15140f;--card:#1d1b15;--ink:#efe8da;--hair:#39342a;--muted:#b5ab98;--forest:#A8D88A;--mid:#8fc775}
*{box-sizing:border-box}html{-webkit-text-size-adjust:100%}
body{margin:0;background:var(--paper);color:var(--ink);font:17px/1.55 system-ui,-apple-system,"Segoe UI",Roboto,Arial,sans-serif}
header{position:sticky;top:0;z-index:5;background:var(--forest);color:#F7F2E9;padding:10px 16px;display:flex;gap:10px;align-items:center}
header b{font-family:Georgia,serif;font-size:17px;flex:1;white-space:nowrap;overflow:hidden;text-overflow:ellipsis}
header select{max-width:52vw;font-size:15px;padding:6px;border-radius:6px;border:0;background:#F7F2E9;color:#20190F}
main{max-width:760px;margin:0 auto;padding:8px 16px 80px}
section.doc{border-bottom:3px solid var(--hair);padding-bottom:24px;margin-bottom:24px}
h1,h2,h3{font-family:Georgia,"Times New Roman",serif;line-height:1.2;color:var(--forest)}
h1{font-size:1.7rem;margin:24px 0 8px}h2{font-size:1.3rem;margin:28px 0 8px;padding-top:12px;border-top:1px solid var(--hair)}
h3{font-size:1.05rem;margin:20px 0 6px;color:var(--ochre)}
table{border-collapse:collapse;width:100%;font-size:14.5px;margin:10px 0;display:block;overflow-x:auto}
th,td{border:1px solid var(--hair);padding:6px 8px;text-align:left;vertical-align:top}
th{background:var(--card)}
code,pre{font-family:ui-monospace,Menlo,monospace;font-size:13.5px;background:var(--card);border-radius:4px}
pre{padding:10px;overflow-x:auto}
blockquote{margin:12px 0;padding:6px 14px;border-left:4px solid var(--ochre);background:var(--card)}
a{color:var(--water)}strong{color:inherit}
li{margin:3px 0}
.meta{color:var(--muted);font-size:13px}
"""


def md_to_html(text: str) -> str:
    import markdown
    return markdown.markdown(text, extensions=["tables", "fenced_code", "sane_lists"])


def build_pack(cdir: Path, out: Path) -> None:
    order = ["README.md", "day-0-preparation.md"] + sorted(p.name for p in cdir.glob("day-*.md")
                                                              if p.name != "day-0-preparation.md")
    rest = sorted(p.name for p in cdir.glob("*.md") if p.name not in order)
    # put the parsed technical files last; they are for producers more than facilitators
    tail = [n for n in ("slides.md", "media.md") if n in rest]
    rest = [n for n in rest if n not in tail] + tail
    docs = [(n, cdir / n) for n in order + rest if (cdir / n).exists()]
    docs += [(f"shared/{n}", COURSES / "shared" / n) for n in ("facilitation-toolkit.md", "forms.md", "poster-standards.md")]
    title = COURSE_TITLES.get(cdir.name, cdir.name)
    opts, bodies = [], []
    for i, (name, path) in enumerate(docs):
        anchor = f"doc{i}"
        src = path.read_text(encoding="utf-8")
        first = next((l[2:].strip() for l in src.splitlines() if l.startswith("# ")), name)
        opts.append(f'<option value="#{anchor}">{html.escape(first)}</option>')
        bodies.append(f'<section class="doc" id="{anchor}"><p class="meta">{html.escape(name)}</p>{md_to_html(src)}</section>')
    page = f"""<!doctype html><html lang="en"><head><meta charset="utf-8">
<meta name="viewport" content="width=device-width,initial-scale=1">
<title>{html.escape(title)} — facilitator pack</title><style>{PACK_CSS}</style></head><body>
<header><b>{html.escape(title)}</b><select aria-label="Jump to" onchange="location.hash=this.value">{''.join(opts)}</select></header>
<main>{''.join(bodies)}</main></body></html>"""
    out.parent.mkdir(parents=True, exist_ok=True)
    out.write_text(page, encoding="utf-8")


# ── generation prompts (for Gemini/Imagen, Veo, Google Flow, ChatGPT) ────────────────────────

FORMATS = {
    "image": "Landscape 16:9 composition.",
    "poster-art": "Portrait 3:4 composition for an A1 teaching poster; keep the top 15% and bottom 20% "
                  "calm and uncluttered (text is added there later).",
    "photo": "Landscape 16:9.",
    "animation": "16:9, 8-second shot, silent.",
}


def attach_prompts(catalog: dict) -> None:
    styles: dict[str, dict] = {}
    for it in catalog.values():
        if it["kind"] == "poster" or it.get("status") != "NEW":
            continue
        st = styles.setdefault(it["course"], load_style(it["course"]))
        anchor = st["animation_anchor"] if it["kind"] == "animation" else st["image_anchor"]
        fmt = FORMATS.get(it["kind"], "")
        if it["kind"] == "image" and "card" in (it.get("length") or ""):
            fmt = "Portrait 2:3 card composition, one subject centred, plain background."
        brief = re.sub(r"\s+", " ", it.get("brief") or it["title"]).strip()
        it["prompt"] = f"{anchor}\n\n{fmt}\n\nSubject: {brief}".strip()
        if it["kind"] == "animation":
            secs = re.search(r"(\d+)\s*s", it.get("length") or "")
            n = max(1, -(-int(secs.group(1)) // 8)) if secs else 1
            it["shots"] = n


def attach_translations(catalog: dict, course_dirs: list[Path]) -> None:
    for cdir in course_dirs:
        zu = load_i18n(cdir.name, "zu").get("posters") or {}
        for pid, text in zu.items():
            if pid in catalog:
                catalog[pid]["zu"] = text


def write_prompt_packs(course: str, catalog: dict, out_dir: Path) -> None:
    st = load_style(course)
    items = [it for it in catalog.values() if it["course"] == course and it.get("prompt")]
    imgs = [it for it in items if it["kind"] != "animation"]
    anis = [it for it in items if it["kind"] == "animation"]
    head = (f"# {COURSE_TITLES.get(course, course)} — %s prompts ({st['look']})\n\n"
            "Generated by scripts/courses/build.py. Paste each block as-is. Save the result at the path shown, "
            "then re-run the build — the decks pick it up.\n\n")
    lines = [head % "image"]
    lines.append("Tips for ChatGPT / Gemini: keep one chat per course so the look stays consistent; generate the "
                 "character sheet first and attach it as a reference to later prompts.\n\n")
    for it in sorted(imgs, key=lambda i: i["id"]):
        lines.append(f"## {it['id']} · {it['title']}\nSave as: `{it['target']}`  ·  used in: {', '.join(it['used_in'])}\n\n```\n{it['prompt']}\n```\n\n")
    (out_dir / "prompts-images.md").write_text("".join(lines), encoding="utf-8")
    lines = [head % "animation"]
    lines.append("Google Flow makes ~8-second shots. Items marked with more than one shot: make shot 1 from the "
                 "prompt, then use Scenebuilder → Extend (or Frames to Video from the last frame) for the rest, "
                 "describing the next step of the Subject. Export MP4 720p, mute the audio track, and save a "
                 "still frame as .jpg with the same name.\n\n")
    for it in sorted(anis, key=lambda i: i["id"]):
        lines.append(f"## {it['id']} · {it['title']}\nSave as: `{it['target']}` (+ `.jpg` still)  ·  shots: {it.get('shots', 1)}  ·  "
                     f"length: {it.get('length', '')}\n\n```\n{it['prompt']}\n```\n\n")
    (out_dir / "prompts-animations.md").write_text("".join(lines), encoding="utf-8")


def _labels(text: str) -> list[str]:
    out = []
    for line in (text or "").splitlines():
        m = re.match(r"^\s*\d+[.)]\s*(.+)$", line)
        if m:
            out.append(m.group(1).strip())
    return out


def _question(layout: str) -> str:
    m = re.search(r"[Qq]uestion strip:\s*\*?\"?([^\"*]+?)\"?\*?(?:\.|$)", layout or "")
    return m.group(1).strip() if m else ""


def write_poster_prompts(course: str, catalog: dict, out_dir: Path) -> None:
    """Full-poster prompts (art + isiZulu/English lettering) for ChatGPT's image mode."""
    st = load_style(course)
    posters = sorted((it for it in catalog.values() if it["course"] == course and it["kind"] == "poster"),
                     key=lambda i: i["id"])
    lines = [f"# {COURSE_TITLES.get(course, course)} — poster prompts, isiZulu + English ({st['look']})\n\n",
             "Generated by scripts/courses/build.py. One block = one finished poster. Paste into ChatGPT "
             "(image mode), one chat per course, portrait. Then CHECK EVERY WORD against the block — image "
             "models misspell, especially isiZulu. If a word is wrong, reply in the same chat: "
             "'Keep everything the same, only correct the spelling of <wrong> to <right>.' "
             "Save the result as the phone JPG path shown; for print, upscale to A1 or rebuild the "
             "lettering in a layout tool from the same words. isiZulu is a DRAFT until read back by a "
             "first-language speaker.\n\n"]
    for it in posters:
        en_labels = _labels(it.get("labels"))
        zu = it.get("zu") or {}
        zu_labels = zu.get("labels") or []
        en_q = _question(it.get("layout"))
        art = catalog.get(it.get("art"), {}).get("brief", "")
        text = [f'- Small top corner: "{it["id"]}"',
                f'- HEADLINE in isiZulu (largest words on the poster): "{zu.get("headline") or it["headline"]}"',
                f'- Directly under it, smaller (about 60%), in English: "{it["headline"]}"']
        if en_labels:
            text.append("- Numbered labels. Each label: isiZulu in bold, English underneath in smaller regular type:")
            for n, en in enumerate(en_labels):
                zl = zu_labels[n] if n < len(zu_labels) else ""
                text.append(f'  {n + 1}. "{zl}" / "{en}"' if zl else f'  {n + 1}. "{en}"')
        q = zu.get("question") or ""
        if q or en_q:
            text.append(f'- Bottom strip, italic: "{q or en_q}"' + (f' (English, smaller: "{en_q}")' if q and en_q else ""))
        layout = " ".join((it.get("layout") or "").split())
        art_flat = " ".join(art.split())
        prompt = (st["poster_anchor"] + "\n\nPortrait 3:4 (A1 poster). Layout: " + layout + "\n\n"
                  + "PICTURE: " + art_flat + "\n\n"
                  + "TEXT - use exactly these words, spelled exactly as written, and no other words, letters, "
                  + "numbers or watermarks anywhere:\n" + "\n".join(text) + "\n\n"
                  + "Text must be large, high-contrast and readable from 6 metres; keep it inside the margins; "
                  + "the poster must still read when photocopied in black and white. SOUTHERN HEMISPHERE: sun in "
                  + "the north, shadows fall south. People are Black South Africans of mixed ages and genders.")
        lines.append(f"## {it['id']} · {it['title']}\nSave as: `{it['phone_target']}`  ·  size: {it.get('size')}  ·  "
                     f"used in: {', '.join(it['used_in'])}\n\n```\n{prompt}\n```\n\n")
    (out_dir / "prompts-posters.md").write_text("".join(lines), encoding="utf-8")


# ── main ─────────────────────────────────────────────────────────────────────────────────────

def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("courses", nargs="*", help="course folder names (default: all)")
    ap.add_argument("--check", action="store_true", help="validate only")
    ap.add_argument("--embed-video", action="store_true", help="embed clips in the .pptx (large files)")
    args = ap.parse_args()

    all_dirs = sorted(p for p in COURSES.iterdir() if p.is_dir() and (p / "slides.md").exists())
    wanted = [d for d in all_dirs if not args.courses or d.name in args.courses]
    if args.courses and len(wanted) != len(args.courses):
        print("Unknown course:", set(args.courses) - {d.name for d in wanted}, file=sys.stderr)
        return 2

    catalog = build_catalog(all_dirs)
    attach_prompts(catalog)
    attach_translations(catalog, all_dirs)
    problems = validate(wanted, catalog)
    for p in problems:
        print("WARN", p)
    if args.check:
        print(f"{len(problems)} problem(s) in {len(wanted)} course(s)")
        return 1 if problems else 0

    manifest = {
        "about": "Generated by scripts/courses/build.py from courses/*/media.md and posters.md. Do not edit by hand.",
        "rules": "courses/shared/poster-standards.md (Art rules) and courses/shared/PRODUCTION-GUIDE.md",
        "items": sorted(catalog.values(), key=lambda it: (it["course"], it["kind"], it["id"])),
    }
    (COURSES / "media-manifest.json").write_text(json.dumps(manifest, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")

    for cdir in wanted:
        decks = parse_slides(cdir / "slides.md")
        langs = ["en"] + [l for l in ("zu",) if (cdir / "i18n" / f"{l}.json").exists()]
        for d in decks:
            for lang in langs:
                name = f"{d.id}.pptx" if lang == "en" else f"{d.id}.{lang}.pptx"
                build_deck(cdir.name, d, catalog, BUILD / cdir.name / name, args.embed_video, lang)
        write_prompt_packs(cdir.name, catalog, BUILD / cdir.name)
        write_poster_prompts(cdir.name, catalog, BUILD / cdir.name)
        build_pack(cdir, BUILD / cdir.name / f"{cdir.name}-phone-pack.html")
        n_new = sum(1 for it in catalog.values() if it["course"] == cdir.name and it["status"] == "NEW" and it["kind"] != "poster")
        n_reuse = sum(1 for it in catalog.values() if it["course"] == cdir.name and it["status"] == "REUSE")
        print(f"{cdir.name} [{'+'.join(langs)}]: {len(decks)} decks, {sum(len(d.slides) for d in decks)} slides; media NEW {n_new} / REUSE {n_reuse}")
    print(f"manifest: {len(catalog)} items → courses/media-manifest.json")
    return 0


if __name__ == "__main__":
    sys.exit(main())
